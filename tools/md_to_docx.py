#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 report/ 下的 Markdown 设计报告渲染为提交用 .docx（python-docx）。

支持：标题（#/##/###）、段落、无序/有序列表、引用、代码块、管道表、图片、
**粗体** 与 `行内代码` 的基础渲染。图片路径相对 Markdown 文件解析。

用法::

    python tools/md_to_docx.py --src report/design_report.md --out report/design_report.docx
"""
import argparse
import re
import sys
from pathlib import Path

sys.dont_write_bytecode = True

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CJK_FONT = "微软雅黑"
MONO_FONT = "Consolas"


def set_run_font(run, name=CJK_FONT, size=None, bold=None, color=None, mono=False):
    run.font.name = MONO_FONT if mono else name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_inline(par, text, size=11, base_bold=False):
    """渲染 **粗体**、`代码`、[文本](链接) → 文本。"""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    tokens = re.split(r"(\*\*.+?\*\*|`.+?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2])
            set_run_font(r, size=size, bold=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            set_run_font(r, name=MONO_FONT, size=size - 0.5, mono=True)
        else:
            r = par.add_run(tok)
            set_run_font(r, size=size, bold=base_bold)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            txt = row[j] if j < len(row) else ""
            add_inline(p, txt.strip(), size=9)
            if i == 0:
                shade(cell, "E8EEF7")
                for r in p.runs:
                    r.bold = True
    doc.add_paragraph()


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(4)
    for i, line in enumerate(lines):
        r = p.add_run(line)
        set_run_font(r, name=MONO_FONT, size=8.5, mono=True)
        if i != len(lines) - 1:
            r.add_break()


def parse_table_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def convert(src: Path, out: Path):
    text = src.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = CJK_FONT
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    for s in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[s]
        st.font.name = CJK_FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    i = 0
    in_code = False
    code_buf = []
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.rstrip()

        if s.strip().startswith("```"):
            if in_code:
                add_code(doc, code_buf)
                code_buf, in_code = [], False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not s.strip():
            i += 1
            continue
        if s.strip() in ("---", "***", "___"):
            i += 1
            continue

        # 图片
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s.strip())
        if m:
            alt, rel = m.group(1), m.group(2)
            img = (src.parent / rel).resolve()
            if img.is_file():
                doc.add_picture(str(img), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(alt)
                set_run_font(r, size=9, color=(0x55, 0x65, 0x77))
            else:
                p = doc.add_paragraph()
                r = p.add_run(f"[缺图] {rel}")
                set_run_font(r, size=9, color=(0xC0, 0x00, 0x00))
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            level, title = len(m.group(1)), m.group(2).strip()
            if level == 1:
                p = doc.add_heading(level=0)
                r = p.add_run(title)
                set_run_font(r, size=18, bold=True)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_heading(level=min(level - 1, 3))
                r = p.add_run(title)
                set_run_font(r, size=15 if level == 2 else 12.5, bold=True)
            i += 1
            continue

        # 表格
        if s.strip().startswith("|") and i + 1 < n and re.match(r"^\s*\|?\s*:?-{3,}", lines[i + 1]):
            rows = [parse_table_row(s)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # 引用
        if s.strip().startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]).strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            add_inline(p, " ".join(buf), size=9.5)
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x65, 0x77)
            continue

        # 列表
        m = re.match(r"^\s*([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            style_name = "List Bullet" if m.group(1) in "-*+" else "List Number"
            p = doc.add_paragraph(style=style_name)
            add_inline(p, m.group(2), size=10.5)
            i += 1
            continue

        # 普通段落（合并连续行）
        buf = [s.strip()]
        i += 1
        while i < n:
            nxt = lines[i].rstrip()
            if (not nxt.strip() or nxt.strip().startswith(("#", "|", ">", "```", "!["))
                    or re.match(r"^\s*([-*+]|\d+\.)\s+", nxt)):
                break
            buf.append(nxt.strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.0)
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, " ".join(buf), size=10.5)

    if in_code and code_buf:
        add_code(doc, code_buf)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"WROTE {out} {out.stat().st_size} bytes")


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=str(Path("report") / "design_report.md"))
    ap.add_argument("--out", default=str(Path("report") / "design_report.docx"))
    args = ap.parse_args(argv)
    src = Path(args.src)
    if not src.is_absolute():
        src = Path.cwd() / src
    out = Path(args.out)
    if not out.is_absolute():
        out = Path.cwd() / out
    if not src.is_file():
        print(f"FAIL src not found: {src}")
        return 1
    convert(src, out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
